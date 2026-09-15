"""
Tests for file format handlers.

Tests extraction of text from .txt, .docx, .xlsx, .pdf, .json files.
"""

import json
import io
import pytest
from src.sme.file_handlers import (
    extract_text_from_file,
    get_handler,
    get_supported_formats,
    FileHandlerError,
    PlainTextHandler,
    DocxHandler,
    ExcelHandler,
    PdfHandler,
    JsonHandler,
    HAS_DOCX,
    HAS_OPENPYXL,
    HAS_PYPDF2,
)


# ---------------------------------------------------------------------------
# PlainTextHandler Tests
# ---------------------------------------------------------------------------

def test_plain_text_handler_can_handle():
    handler = PlainTextHandler()
    assert handler.can_handle("response.txt")
    assert handler.can_handle("file.TXT")
    assert not handler.can_handle("file.docx")


def test_plain_text_handler_extract_utf8():
    handler = PlainTextHandler()
    content = "This is a test.\nLine 2.".encode('utf-8')
    result = handler.extract_text(content)
    assert result == "This is a test.\nLine 2."


def test_plain_text_handler_extract_latin1():
    handler = PlainTextHandler()
    # Invalid UTF-8 sequence, should fallback to latin-1
    content = b"Test \xe9 accented"
    result = handler.extract_text(content)
    assert "Test" in result


# ---------------------------------------------------------------------------
# JsonHandler Tests
# ---------------------------------------------------------------------------

def test_json_handler_can_handle():
    handler = JsonHandler()
    assert handler.can_handle("responses.json")
    assert handler.can_handle("file.JSON")
    assert not handler.can_handle("file.txt")


def test_json_handler_extract_valid():
    handler = JsonHandler()
    data = {"gap_id": "gap-001", "answer": "test answer"}
    content = json.dumps(data).encode('utf-8')
    result = handler.extract_text(content)
    assert "gap_id" in result
    assert "gap-001" in result


def test_json_handler_extract_invalid():
    handler = JsonHandler()
    content = b"{invalid json"
    with pytest.raises(FileHandlerError, match="Invalid JSON"):
        handler.extract_text(content)


# ---------------------------------------------------------------------------
# DocxHandler Tests (conditional on python-docx)
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_handler_can_handle():
    handler = DocxHandler()
    assert handler.can_handle("response.docx")
    assert handler.can_handle("file.DOCX")
    assert not handler.can_handle("file.txt")


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_handler_extract_with_paragraphs():
    handler = DocxHandler()
    from docx import Document
    
    # Create a simple DOCX in memory
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    
    bytes_io = io.BytesIO()
    doc.save(bytes_io)
    content = bytes_io.getvalue()
    
    result = handler.extract_text(content)
    assert "First paragraph" in result
    assert "Second paragraph" in result


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_handler_extract_with_table():
    handler = DocxHandler()
    from docx import Document
    
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header1"
    table.rows[0].cells[1].text = "Header2"
    table.rows[1].cells[0].text = "Value1"
    table.rows[1].cells[1].text = "Value2"
    
    bytes_io = io.BytesIO()
    doc.save(bytes_io)
    content = bytes_io.getvalue()
    
    result = handler.extract_text(content)
    assert "Header1" in result
    assert "Value1" in result


def test_docx_handler_missing_library():
    """Test that DocxHandler raises clear error if python-docx not available."""
    handler = DocxHandler()
    if not HAS_DOCX:
        with pytest.raises(FileHandlerError, match="python-docx not installed"):
            handler.extract_text(b"dummy")


# ---------------------------------------------------------------------------
# ExcelHandler Tests (conditional on openpyxl)
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
def test_excel_handler_can_handle():
    handler = ExcelHandler()
    assert handler.can_handle("response.xlsx")
    assert handler.can_handle("file.XLSX")
    assert handler.can_handle("file.xls")
    assert not handler.can_handle("file.txt")


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
def test_excel_handler_extract_single_sheet():
    handler = ExcelHandler()
    from openpyxl import Workbook
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws['A1'] = "Header"
    ws['A2'] = "Value"
    
    bytes_io = io.BytesIO()
    wb.save(bytes_io)
    content = bytes_io.getvalue()
    
    result = handler.extract_text(content)
    assert "Sheet1" in result
    assert "Header" in result
    assert "Value" in result


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
def test_excel_handler_extract_multiple_sheets():
    handler = ExcelHandler()
    from openpyxl import Workbook
    
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1['A1'] = "Sheet1Data"
    
    ws2 = wb.create_sheet("Sheet2")
    ws2['A1'] = "Sheet2Data"
    
    bytes_io = io.BytesIO()
    wb.save(bytes_io)
    content = bytes_io.getvalue()
    
    result = handler.extract_text(content)
    assert "Sheet1" in result
    assert "Sheet2" in result
    assert "Sheet1Data" in result
    assert "Sheet2Data" in result


def test_excel_handler_missing_library():
    """Test that ExcelHandler raises clear error if openpyxl not available."""
    handler = ExcelHandler()
    if not HAS_OPENPYXL:
        with pytest.raises(FileHandlerError, match="openpyxl not installed"):
            handler.extract_text(b"dummy")


# ---------------------------------------------------------------------------
# PdfHandler Tests (conditional on PyPDF2)
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not HAS_PYPDF2, reason="PyPDF2 not installed")
def test_pdf_handler_can_handle():
    handler = PdfHandler()
    assert handler.can_handle("response.pdf")
    assert handler.can_handle("file.PDF")
    assert not handler.can_handle("file.txt")


def test_pdf_handler_missing_library():
    """Test that PdfHandler raises clear error if PyPDF2 not available."""
    handler = PdfHandler()
    if not HAS_PYPDF2:
        with pytest.raises(FileHandlerError, match="PyPDF2 not installed"):
            handler.extract_text(b"dummy")


# ---------------------------------------------------------------------------
# get_handler() Tests
# ---------------------------------------------------------------------------

def test_get_handler_txt():
    handler = get_handler("file.txt")
    assert handler is not None
    assert isinstance(handler, PlainTextHandler)


def test_get_handler_json():
    handler = get_handler("file.json")
    assert handler is not None
    assert isinstance(handler, JsonHandler)


def test_get_handler_docx():
    handler = get_handler("file.docx")
    if HAS_DOCX:
        assert handler is not None
        assert isinstance(handler, DocxHandler)
    else:
        # docx handler still registered, but will fail at extraction
        assert handler is not None


def test_get_handler_xlsx():
    handler = get_handler("file.xlsx")
    if HAS_OPENPYXL:
        assert handler is not None
        assert isinstance(handler, ExcelHandler)
    else:
        assert handler is not None


def test_get_handler_unsupported():
    handler = get_handler("file.xyz")
    assert handler is None


# ---------------------------------------------------------------------------
# extract_text_from_file() Tests
# ---------------------------------------------------------------------------

def test_extract_text_from_file_txt():
    content = "Test content".encode('utf-8')
    result = extract_text_from_file("response.txt", content)
    assert result == "Test content"


def test_extract_text_from_file_json():
    data = {"key": "value"}
    content = json.dumps(data).encode('utf-8')
    result = extract_text_from_file("response.json", content)
    assert "key" in result
    assert "value" in result


def test_extract_text_from_file_unsupported():
    with pytest.raises(FileHandlerError, match="Unsupported file type"):
        extract_text_from_file("file.xyz", b"content")


def test_extract_text_from_file_case_insensitive():
    content = "Test".encode('utf-8')
    result = extract_text_from_file("FILE.TXT", content)
    assert result == "Test"


# ---------------------------------------------------------------------------
# get_supported_formats() Tests
# ---------------------------------------------------------------------------

def test_get_supported_formats():
    formats = get_supported_formats()
    assert ".txt" in formats
    assert ".json" in formats
    assert all(f.startswith('.') for f in formats)
    assert len(formats) >= 2  # At least .txt and .json
    
    if HAS_DOCX:
        assert ".docx" in formats
    if HAS_OPENPYXL:
        assert ".xlsx" in formats
    if HAS_PYPDF2:
        assert ".pdf" in formats


# ---------------------------------------------------------------------------
# Integration Tests
# ---------------------------------------------------------------------------

def test_extract_with_ref_pattern_in_text():
    """Verify text with [REF:...] patterns is extracted correctly."""
    content = b"[REF:gap-001]: This is the answer\n[REF:gap-002]: Another answer"
    result = extract_text_from_file("response.txt", content)
    assert "[REF:gap-001]" in result
    assert "This is the answer" in result


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_with_ref_pattern():
    """Verify DOCX extraction preserves [REF:...] patterns."""
    from docx import Document
    
    doc = Document()
    doc.add_paragraph("[REF:gap-001]: Source is sales_db.customer_transactions")
    doc.add_paragraph("[REF:gap-002]: Target is adls://lakehouse/bronze/customer")
    
    bytes_io = io.BytesIO()
    doc.save(bytes_io)
    content = bytes_io.getvalue()
    
    result = extract_text_from_file("response.docx", content)
    assert "[REF:gap-001]" in result
    assert "sales_db.customer_transactions" in result
    assert "[REF:gap-002]" in result
    assert "adls://lakehouse/bronze/customer" in result
